"""XLSX table loader using openpyxl in read-only mode."""

from itertools import islice
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from rag_kit.domain.errors import IngestionError
from rag_kit.domain.models import RawTable


class XlsxLoader:
    def load(self, path: Path) -> RawTable:
        try:
            workbook = load_workbook(path, read_only=True, data_only=True)
            sheet = next(
                candidate
                for candidate in workbook.worksheets
                if candidate.max_row and candidate.max_column
            )
            values = iter(sheet.iter_rows(values_only=True))
            header_candidates = list(islice(values, 10))
            header_index = max(
                range(len(header_candidates)),
                key=lambda index: sum(value is not None for value in header_candidates[index]),
            )
            names = self._headers(header_candidates[header_index])
            rows = tuple(
                dict(zip(names, row, strict=False))
                for row in (*header_candidates[header_index + 1 :], *values)
                if any(value is not None for value in row)
            )
            sheet_name = sheet.title
            workbook.close()
        except (OSError, ValueError, StopIteration, IndexError) as exc:
            raise IngestionError(f"could not read XLSX {path}: {exc}") from exc  # noqa: TRY003
        normalized: tuple[dict[str, Any], ...] = tuple(dict(row) for row in rows)
        return RawTable(
            columns=tuple(names),
            source_name=str(path),
            rows=normalized,
            warnings=(f"sheet: {sheet_name}; header row: {header_index + 1}",),
        )

    def _headers(self, headers: tuple[Any, ...]) -> list[str]:
        names: list[str] = []
        seen: dict[str, int] = {}
        for index, value in enumerate(headers):
            name = str(value).strip() if value is not None else f"column_{index}"
            seen[name] = seen.get(name, 0) + 1
            names.append(f"{name}_{seen[name]}" if seen[name] > 1 else name)
        return names
